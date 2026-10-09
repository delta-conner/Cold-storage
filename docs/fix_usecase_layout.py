# -*- coding: utf-8 -*-
"""按范本「图3.5」风格重绘用例图：少线、不挡字、椭圆全在框内。"""
from pathlib import Path
import matplotlib.pyplot as plt
from matplotlib.patches import Ellipse, Rectangle, FancyArrowPatch
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(r"f:\plcmin")
FIG = ROOT / "docs" / "thesis_figs"
FIG.mkdir(parents=True, exist_ok=True)
DOCS = [
    ROOT / "docs" / "冷库设备运维管理系统-论文第1-4章-用例图已改样式.docx",
    ROOT / "docs" / "冷库设备运维管理系统-论文第1-4章.docx",
    ROOT / "docs" / "冷库设备运维管理系统-论文前三章-用例图已改样式.docx",
    ROOT / "docs" / "冷库设备运维管理系统-论文前三章.docx",
]

plt.rcParams["font.sans-serif"] = ["SimSun", "Microsoft YaHei", "SimHei"]
plt.rcParams["axes.unicode_minus"] = False


def save(fig, name):
    p = FIG / name
    fig.savefig(p, dpi=200, bbox_inches="tight", facecolor="white", pad_inches=0.15)
    plt.close(fig)
    return p


def actor(ax, x, y, label):
    ax.add_patch(plt.Circle((x, y), 0.12, fill=False, ec="black", lw=1.35, zorder=5))
    ax.plot([x, x], [y - 0.12, y - 0.48], "k-", lw=1.35, zorder=5)
    ax.plot([x - 0.18, x + 0.18], [y - 0.25, y - 0.25], "k-", lw=1.35, zorder=5)
    ax.plot([x, x - 0.15], [y - 0.48, y - 0.72], "k-", lw=1.35, zorder=5)
    ax.plot([x, x + 0.15], [y - 0.48, y - 0.72], "k-", lw=1.35, zorder=5)
    ax.text(x, y - 0.95, label, ha="center", va="top", fontsize=12)


def uc(ax, x, y, text, w=2.35, h=0.78):
    ax.add_patch(Ellipse((x, y), w, h, facecolor="white", edgecolor="black", lw=1.25, zorder=4))
    ax.text(x, y, text, ha="center", va="center", fontsize=11, zorder=6)
    return (x, y, w, h)


def boundary(ax, x0, y0, x1, y1, title):
    ax.add_patch(Rectangle((x0, y0), x1 - x0, y1 - y0, fill=False, ec="black", lw=1.5, zorder=1))
    # 标题留足上方空白，避免被椭圆挡住
    ax.text((x0 + x1) / 2, y1 - 0.35, title, ha="center", va="top", fontsize=13, fontweight="bold")


def solid(ax, x1, y1, x2, y2):
    ax.plot([x1, x2], [y1, y2], "k-", lw=1.05, zorder=2)


def include(ax, x1, y1, x2, y2, label_side=0.22):
    """虚线 include，标签放在线旁，避免压在线上。"""
    ax.annotate(
        "",
        xy=(x2, y2),
        xytext=(x1, y1),
        arrowprops=dict(arrowstyle="->", color="black", lw=1.0, linestyle=(0, (4, 3))),
        zorder=3,
    )
    mx, my = (x1 + x2) / 2, (y1 + y2) / 2
    # 根据方向把标签偏到一侧
    dx, dy = x2 - x1, y2 - y1
    # 法向偏移
    import math
    L = math.hypot(dx, dy) or 1
    nx, ny = -dy / L * label_side, dx / L * label_side
    ax.text(mx + nx, my + ny, "<<include>>", ha="center", va="center", fontsize=9, zorder=7,
            bbox=dict(boxstyle="round,pad=0.12", facecolor="white", edgecolor="none", alpha=0.92))


def fig_detail(fname, title, actor_name, left_ucs, right_top, right_bottom_login=True):
    """
    仿图3.5：左列主用例（与参与者相连），右列辅助用例，登录在右下，include 干净。
    left_ucs: [str, ...] 1~3 个
    right_top: str 可选，被 include 的中间用例；若 None 则 left 直接 include 登录
    """
    fig, ax = plt.subplots(figsize=(9.0, 5.6))
    ax.set_xlim(0, 9.0)
    ax.set_ylim(0, 5.6)
    ax.axis("off")
    boundary(ax, 2.15, 0.45, 8.7, 5.25, title)
    actor(ax, 0.95, 3.05, actor_name)

    # 左列位置
    n = len(left_ucs)
    left_ys = [3.85, 2.35] if n == 2 else ([4.05, 3.05, 2.05] if n == 3 else [3.1])
    left_pos = []
    for y, name in zip(left_ys[:n], left_ucs):
        left_pos.append(uc(ax, 4.0, y, name, w=2.4, h=0.82))
        # 从参与者不同高度出线，避免粘成一束
        solid(ax, 1.15, 2.9 + (y - 3.1) * 0.15, 2.75, y)

    if right_top:
        rt = uc(ax, 6.85, 3.85, right_top, w=2.3, h=0.82)
        for x, y, w, h in left_pos:
            include(ax, x + w * 0.42, y, rt[0] - rt[2] * 0.42, rt[1], label_side=0.28)
        if right_bottom_login:
            lg = uc(ax, 6.85, 2.15, "登录", w=2.0, h=0.78)
            include(ax, rt[0], rt[1] - rt[3] * 0.42, lg[0], lg[1] + lg[3] * 0.42, label_side=0.35)
    elif right_bottom_login:
        lg = uc(ax, 6.85, 3.1, "登录", w=2.0, h=0.78)
        for x, y, w, h in left_pos:
            include(ax, x + w * 0.42, y, lg[0] - lg[2] * 0.42, lg[1], label_side=0.32)

    return save(fig, fname)


def fig_actor_overview(fname, title, actor_name, cases):
    """
    角色总览：两列椭圆，全部在框内；参与者只连左列（或全部但分高出线）；
    不画满屏 include——只让「登录」被一个「身份认证相关」概括，或左列最后一项为登录由参与者直连。
    做法：用例全部直连参与者，不加 include（总览图更清晰）；登录作为普通用例之一放右下。
    """
    # 两列布局
    cols = 2
    n = len(cases)
    rows = (n + 1) // 2
    height = 1.6 + rows * 1.05 + 0.9
    width = 10.2
    fig, ax = plt.subplots(figsize=(width, height * 0.95))
    ax.set_xlim(0, width)
    ax.set_ylim(0, height)
    ax.axis("off")

    x0, y0, x1, y1 = 2.3, 0.4, 9.95, height - 0.25
    boundary(ax, x0, y0, x1, y1, title)
    ay = (y0 + y1) / 2 + 0.2
    actor(ax, 1.05, ay, actor_name)

    # 标题下方可用区域
    top = y1 - 0.85
    bot = y0 + 0.55
    positions = []
    for i, name in enumerate(cases):
        col = i % 2
        row = i // 2
        # 行从上往下
        if rows == 1:
            y = (top + bot) / 2
        else:
            y = top - row * ((top - bot) / max(rows - 1, 1))
        x = 4.15 if col == 0 else 7.55
        positions.append((uc(ax, x, y, name, w=2.55, h=0.78), i, col, y))

    # 参与者连到每一个用例，但从不同锚点出线（沿身体高度分散）
    for (x, y, w, h), i, col, yy in positions:
        # 出线点：在参与者右侧，按索引上下错开
        t = i / max(len(cases) - 1, 1)
        out_y = ay - 0.15 + (t - 0.5) * 1.1
        solid(ax, 1.28, out_y, x - w * 0.48, y)

    return save(fig, fname)


def fig_role():
    """三角色关系总览：无 include，留白充足，标题不被挡。"""
    fig, ax = plt.subplots(figsize=(11.0, 7.2))
    ax.set_xlim(0, 11)
    ax.set_ylim(0, 7.2)
    ax.axis("off")
    boundary(ax, 2.6, 0.4, 10.6, 6.9, "冷库设备综合运维管理系统")

    actor(ax, 1.15, 5.7, "管理员")
    actor(ax, 1.15, 3.7, "运维人员")
    actor(ax, 1.15, 1.7, "甲方用户")

    cases = [
        "用户与权限管理",
        "冷库/设备台账管理",
        "工单全流程管控",
        "维保计划与任务",
        "故障/案例/备件",
        "健康分析与统计",
        "消息中心",
        "登录认证",
    ]
    # 两列 4 行，顶部留空
    xs = [5.0, 8.3]
    top, bot = 6.15, 1.05
    pos = []
    for i, name in enumerate(cases):
        col, row = i % 2, i // 2
        y = top - row * ((top - bot) / 3)
        x = xs[col]
        pos.append((uc(ax, x, y, name, w=2.6, h=0.72), name, col, row, x, y))

    # 管理员：全部；运维：除用户权限；甲方：台账/工单/消息/登录
    admin_idx = range(8)
    ops_idx = [1, 2, 3, 4, 5, 6, 7]
    client_idx = [1, 2, 6, 7]

    def connect(actor_y, indices, spread=0.9):
        for k, i in enumerate(indices):
            _, _, _, _, x, y = pos[i]
            t = k / max(len(indices) - 1, 1)
            oy = actor_y - 0.1 + (t - 0.5) * spread
            solid(ax, 1.38, oy, x - 1.35, y)

    connect(5.55, list(admin_idx), 1.3)
    connect(3.55, list(ops_idx), 1.1)
    connect(1.55, list(client_idx), 0.7)
    return save(fig, "fig3_1_role.png")


def redraw():
    fig_role()

    # 角色总览：无扎堆 include
    fig_actor_overview(
        "fig3_2_client.png", "冷库设备运维管理系统", "甲方用户",
        ["浏览设备台账", "提交报修工单", "查看本人工单", "确认验收",
         "服务评价", "个人中心", "消息中心", "登录"])
    fig_actor_overview(
        "fig3_5_ops.png", "冷库设备运维管理系统", "运维人员",
        ["运维工作台", "处理指派工单", "故障记录", "案例库检索",
         "维保任务执行", "备件领用", "设备健康查看", "消息中心"])
    fig_actor_overview(
        "fig3_7_admin.png", "冷库设备运维管理系统", "管理员",
        ["用户权限管理", "冷库冷藏间管理", "设备台账管理", "工单派单管控",
         "维保计划任务", "备件供应商", "健康分析统计", "操作日志"])

    # 分项图：严格仿图3.5
    fig_detail("fig3_3_device_view.png", "浏览设备台账", "甲方用户",
               ["条件查询设备", "选择报修设备"], "查看设备详情", True)
    fig_detail("fig3_4_submit.png", "提交报修工单", "甲方用户",
               ["填写故障描述", "上传现场图片"], "提交生成工单", True)
    fig_detail("fig3_6_process.png", "处理运维工单", "运维人员",
               ["接单", "填写处理过程", "提交完工"], "登记备件出库", True)
    fig_detail("fig3_8_user_mgmt.png", "管理用户", "管理员",
               ["新增用户", "编辑用户", "启用/禁用"], "查询用户", True)
    fig_detail("fig3_9_device_mgmt.png", "设备台账管理", "管理员",
               ["设备增删改查", "状态变更"], "Excel导入导出", True)
    fig_detail("fig3_10_order_admin.png", "工单全流程管控", "管理员",
               ["派单", "转派", "撤回"], "关闭/归档", True)

    print("figs ok")


def swap_in_docs():
    caption_map = {
        "图3.1": "fig3_1_role.png",
        "图3.2": "fig3_2_client.png",
        "图3.3": "fig3_3_device_view.png",
        "图3.4": "fig3_4_submit.png",
        "图3.5": "fig3_5_ops.png",
        "图3.6": "fig3_6_process.png",
        "图3.7": "fig3_7_admin.png",
        "图3.8": "fig3_8_user_mgmt.png",
        "图3.9": "fig3_9_device_mgmt.png",
        "图3.10": "fig3_10_order_admin.png",
    }
    for path in DOCS:
        if not path.exists():
            continue
        doc = Document(str(path))
        paras = list(doc.paragraphs)
        replaced = 0
        for i, p in enumerate(paras):
            text = (p.text or "").replace(" ", "").replace("　", "")
            hit = None
            for key in sorted(caption_map, key=len, reverse=True):
                if text.startswith(key.replace(" ", "")):
                    hit = caption_map[key]
                    break
            if not hit:
                continue
            img = FIG / hit
            j = i - 1
            while j >= 0 and not paras[j]._p.xpath(".//*[local-name()='drawing']"):
                j -= 1
            if j < 0:
                continue
            prev = paras[j]
            for child in list(prev._p):
                prev._p.remove(child)
            prev.add_run().add_picture(str(img), width=Cm(13.5))
            prev.alignment = WD_ALIGN_PARAGRAPH.CENTER
            replaced += 1
        out = path.with_name(path.stem.replace("-用例图已改样式", "") + "-用例图排版优化.docx")
        try:
            doc.save(str(path))
            print("updated", path.name, replaced)
        except PermissionError:
            doc.save(str(out))
            print("locked, wrote", out.name, replaced)


if __name__ == "__main__":
    redraw()
    swap_in_docs()
