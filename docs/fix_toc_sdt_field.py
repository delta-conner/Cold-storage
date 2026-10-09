# -*- coding: utf-8 -*-
"""
把目录改成与周思汉范文一样的「可点击整块」自动目录：
SDT 内容控件 + TOC 域（点一下选中整块，出现更新目录）。
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FINAL = Path(r"f:\plcmin\docs\冷库设备运维管理系统-论文第1-4章.docx")


def set_run_font(run, name_cn="宋体", name_en="Times New Roman", size=12, bold=False):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name_en
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name_cn)


def delete_paragraph(p):
    el = p._element
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def find_toc_title(doc):
    for i, p in enumerate(doc.paragraphs):
        t = (p.text or "").strip().replace(" ", "").replace("\u3000", "")
        if t == "目录":
            return i
    raise SystemExit("未找到目　录")


def find_real_body_ch1(doc, toc_idx):
    for i, p in enumerate(doc.paragraphs):
        if i <= toc_idx:
            continue
        raw = p.text or ""
        if "\t" in raw:
            continue
        t = raw.strip().lstrip("\u3000")
        if not t.startswith("第1章"):
            continue
        for j in range(i + 1, min(i + 12, len(doc.paragraphs))):
            if len((doc.paragraphs[j].text or "").strip()) > 80:
                return i
    raise SystemExit("未找到正文第1章")


def ensure_body_headings(doc, body_idx):
    """保证正文标题仍是 Heading，供 TOC 域抓取。"""
    for i, p in enumerate(doc.paragraphs):
        if i < body_idx:
            continue
        t = (p.text or "").strip().lstrip("\u3000")
        if re.match(r"^第[0-9一二三四五六七八九十]+章", t):
            lvl = 1
        elif re.match(r"^\d+\.\d+\.\d+", t):
            lvl = 3
        elif re.match(r"^\d+\.\d+", t):
            lvl = 2
        else:
            continue
        name = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}[lvl]
        try:
            p.style = doc.styles[name]
        except KeyError:
            continue
        p.paragraph_format.first_line_indent = Cm(0)
        if lvl == 1:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def make_sdt_toc():
    """
    构造与 Word/WPS 自动目录类似的 SDT：
    点一下会选中整块，出现「更新目录」。
    """
    sdt = OxmlElement("w:sdt")
    sdtPr = OxmlElement("w:sdtPr")

    # 文档部件：Table of Contents（WPS/Word 识别为目录块）
    docPartObj = OxmlElement("w:docPartObj")
    docPartGallery = OxmlElement("w:docPartGallery")
    docPartGallery.set(qn("w:val"), "Table of Contents")
    docPartUnique = OxmlElement("w:docPartUnique")
    docPartObj.append(docPartGallery)
    docPartObj.append(docPartUnique)
    sdtPr.append(docPartObj)

    # id
    id_el = OxmlElement("w:id")
    id_el.set(qn("w:val"), "-2067850240")
    sdtPr.append(id_el)

    sdt.append(sdtPr)

    sdtContent = OxmlElement("w:sdtContent")

    # ---- 目录标题段（在 SDT 内，仿范文）----
    # 范文标题有时在 SDT 外；但整块可点通常是 TOC 域在 SDT 内。
    # 这里：SDT 内只放 TOC 域；「目　录」标题保留在外面。

    p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    # 用 TOC 1 样式占位，域更新后会生成各级 TOC 样式
    pStyle = OxmlElement("w:pStyle")
    pStyle.set(qn("w:val"), "TOC1")
    pPr.append(pStyle)
    p.append(pPr)

    def add_run_with(*children):
        r = OxmlElement("w:r")
        for c in children:
            r.append(c)
        p.append(r)

    # begin
    fc1 = OxmlElement("w:fldChar")
    fc1.set(qn("w:fldCharType"), "begin")
    add_run_with(fc1)

    # instr
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    add_run_with(instr)

    # separate
    fc2 = OxmlElement("w:fldChar")
    fc2.set(qn("w:fldCharType"), "separate")
    add_run_with(fc2)

    # 域结果占位（更新后会被替换成真实目录）
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = "请右键此处，选择“更新域”，再选“更新整个目录”。"
    rPr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "21")
    rPr.append(sz)
    r = OxmlElement("w:r")
    r.append(rPr)
    r.append(t)
    p.append(r)

    # end
    fc3 = OxmlElement("w:fldChar")
    fc3.set(qn("w:fldCharType"), "end")
    add_run_with(fc3)

    sdtContent.append(p)
    sdt.append(sdtContent)
    return sdt


def restyle_toc_title(p, doc):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    for child in list(p._element):
        if child.tag.endswith("}r"):
            p._element.remove(child)
    try:
        p.style = doc.styles["Normal"]
    except Exception:
        pass
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(18)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run("目　录")
    set_run_font(run, name_cn="黑体", size=16, bold=True)


def main():
    doc = Document(str(FINAL))
    toc_idx = find_toc_title(doc)
    body_idx = find_real_body_ch1(doc, toc_idx)
    print("toc", toc_idx, "body", body_idx)

    # 删掉目录标题与正文之间一切旧内容（静态条目/旧域）
    for i in range(body_idx - 1, toc_idx, -1):
        delete_paragraph(doc.paragraphs[i])

    toc_idx = find_toc_title(doc)
    body_idx = find_real_body_ch1(doc, toc_idx)
    ensure_body_headings(doc, body_idx)

    toc_title = doc.paragraphs[toc_idx]
    restyle_toc_title(toc_title, doc)

    # 在「目　录」后插入 SDT+TOC 域
    sdt = make_sdt_toc()
    toc_title._element.addnext(sdt)

    # 再插一行操作说明（SDT 外，普通文字）
    tip = OxmlElement("w:p")
    sdt.addnext(tip)
    from docx.text.paragraph import Paragraph

    tip_p = Paragraph(tip, toc_title._parent)
    tip_p.paragraph_format.space_before = Pt(8)
    tip_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = tip_p.add_run(
        "操作：用 WPS/Word 打开后，单击上方灰色目录块 → 点「更新目录」；"
        "或右键目录 → 更新域 → 更新整个目录。"
    )
    set_run_font(r, size=10.5)

    try:
        doc.save(str(FINAL))
        out = FINAL
    except PermissionError:
        out = FINAL.with_name("冷库设备运维管理系统-论文第1-4章-自动目录.docx")
        doc.save(str(out))
    print("saved", out.name)


if __name__ == "__main__":
    main()
