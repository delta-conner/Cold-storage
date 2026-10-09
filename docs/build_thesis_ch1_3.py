# -*- coding: utf-8 -*-
"""生成毕设论文：封面~第3章需求分析（含活动图/用例图）"""
from pathlib import Path
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib.patches import FancyBboxPatch, Ellipse, FancyArrowPatch, Rectangle
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(r"f:\plcmin")
FIG = ROOT / "docs" / "thesis_figs"
FIG.mkdir(parents=True, exist_ok=True)
OUT = ROOT / "docs" / "冷库设备运维管理系统-论文前三章.docx"

plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "SimSun"]
plt.rcParams["axes.unicode_minus"] = False


def savefig(fig, name):
    path = FIG / name
    fig.savefig(path, dpi=160, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def draw_actor(ax, x, y, label, h=0.55):
    ax.plot([x, x], [y - 0.08, y - 0.35], color="#222", lw=1.4)
    ax.plot([x - 0.12, x + 0.12], [y - 0.18, y - 0.18], color="#222", lw=1.4)
    ax.plot([x, x - 0.12], [y - 0.35, y - 0.52], color="#222", lw=1.4)
    ax.plot([x, x + 0.12], [y - 0.35, y - 0.52], color="#222", lw=1.4)
    c = plt.Circle((x, y), 0.08, fill=False, ec="#222", lw=1.4)
    ax.add_patch(c)
    ax.text(x, y - 0.68, label, ha="center", va="top", fontsize=9)


def draw_usecase(ax, x, y, w, h, text):
    e = Ellipse((x, y), w, h, facecolor="#E8F4FC", edgecolor="#1a6b7c", lw=1.2)
    ax.add_patch(e)
    ax.text(x, y, text, ha="center", va="center", fontsize=8, wrap=True)


def draw_system_box(ax, x0, y0, x1, y1, title):
    rect = FancyBboxPatch((x0, y0), x1 - x0, y1 - y0,
                          boxstyle="round,pad=0.02,rounding_size=0.05",
                          facecolor="#FAFBFC", edgecolor="#333", lw=1.3)
    ax.add_patch(rect)
    ax.text((x0 + x1) / 2, y1 - 0.18, title, ha="center", va="top", fontsize=10, fontweight="bold")


def link(ax, x1, y1, x2, y2):
    ax.annotate("", xy=(x2, y2), xytext=(x1, y1),
                arrowprops=dict(arrowstyle="-", color="#555", lw=1.0))


# ---------- figures ----------
def fig_role_usecase():
    fig, ax = plt.subplots(figsize=(10.5, 6.8))
    ax.set_xlim(0, 10.5)
    ax.set_ylim(0, 6.8)
    ax.axis("off")
    draw_system_box(ax, 2.2, 0.4, 8.3, 6.4, "冷库设备综合运维管理系统")
    draw_actor(ax, 0.9, 5.2, "管理员")
    draw_actor(ax, 0.9, 3.4, "运维人员")
    draw_actor(ax, 0.9, 1.6, "甲方用户")
    cases = [
        (5.2, 5.7, "用户与权限管理"),
        (5.2, 5.0, "冷库/设备台账管理"),
        (5.2, 4.3, "工单全流程管控"),
        (5.2, 3.6, "维保计划与任务"),
        (5.2, 2.9, "故障/案例/备件"),
        (5.2, 2.2, "健康分析与统计"),
        (5.2, 1.5, "消息中心"),
        (5.2, 0.85, "登录认证/个人中心"),
    ]
    for x, y, t in cases:
        draw_usecase(ax, x, y, 2.6, 0.55, t)
    # admin links
    for y in [5.7, 5.0, 4.3, 3.6, 2.9, 2.2, 1.5, 0.85]:
        link(ax, 1.15, 5.0, 3.85, y)
    # ops
    for y in [5.0, 4.3, 3.6, 2.9, 2.2, 1.5, 0.85]:
        link(ax, 1.15, 3.2, 3.85, y)
    # client
    for y in [5.0, 4.3, 1.5, 0.85]:
        link(ax, 1.15, 1.4, 3.85, y)
    ax.set_title("图3.1 角色关系用例图", fontsize=12, pad=8)
    return savefig(fig, "fig3_1_role.png")


def fig_actor_usecase(title, fname, actor, cases):
    fig, ax = plt.subplots(figsize=(9.2, 5.8))
    ax.set_xlim(0, 9.2)
    ax.set_ylim(0, 5.8)
    ax.axis("off")
    draw_system_box(ax, 2.4, 0.35, 8.7, 5.45, "系统")
    draw_actor(ax, 1.0, 3.1, actor)
    n = len(cases)
    ys = [4.8 - i * (4.2 / max(n - 1, 1)) for i in range(n)] if n > 1 else [3.0]
    for y, t in zip(ys, cases):
        draw_usecase(ax, 5.5, y, 2.7, 0.52, t)
        link(ax, 1.25, 2.9, 4.1, y)
    ax.set_title(title, fontsize=12, pad=8)
    return savefig(fig, fname)


def fig_activity_order():
    """报修-维修闭环活动图"""
    fig, ax = plt.subplots(figsize=(8.2, 10.2))
    ax.set_xlim(0, 8.2)
    ax.set_ylim(0, 10.2)
    ax.axis("off")

    def node(x, y, w, h, text, kind="action"):
        if kind == "start":
            ax.add_patch(plt.Circle((x, y), 0.12, color="#222"))
        elif kind == "end":
            ax.add_patch(plt.Circle((x, y), 0.14, fill=False, ec="#222", lw=2))
            ax.add_patch(plt.Circle((x, y), 0.07, color="#222"))
        elif kind == "decision":
            diamond = plt.Polygon([(x, y + h / 2), (x + w / 2, y), (x, y - h / 2), (x - w / 2, y)],
                                  closed=True, facecolor="#FFF7E6", edgecolor="#D48806", lw=1.2)
            ax.add_patch(diamond)
            ax.text(x, y, text, ha="center", va="center", fontsize=8)
        else:
            box = FancyBboxPatch((x - w / 2, y - h / 2), w, h,
                                 boxstyle="round,pad=0.02,rounding_size=0.04",
                                 facecolor="#E6F4FF", edgecolor="#1677FF", lw=1.1)
            ax.add_patch(box)
            ax.text(x, y, text, ha="center", va="center", fontsize=8)

    def arr(x1, y1, x2, y2, text=None):
        ax.annotate("", xy=(x2, y2), xytext=(x1, y1),
                    arrowprops=dict(arrowstyle="->", color="#333", lw=1.1))
        if text:
            ax.text((x1 + x2) / 2 + 0.25, (y1 + y2) / 2, text, fontsize=7, color="#666")

    # swimlane-ish vertical flow
    node(4.1, 9.7, 0, 0, "", "start")
    node(4.1, 9.1, 2.6, 0.45, "甲方选择设备提交报修")
    arr(4.1, 9.58, 4.1, 9.35)
    node(4.1, 8.35, 2.6, 0.45, "系统生成待受理工单\n推送管理员消息")
    arr(4.1, 8.85, 4.1, 8.6)
    node(4.1, 7.55, 2.6, 0.45, "管理员派单给运维")
    arr(4.1, 8.1, 4.1, 7.8)
    node(4.1, 6.75, 2.6, 0.45, "运维接单并处理\n填报过程/备件/图片")
    arr(4.1, 7.3, 4.1, 7.0)
    node(4.1, 5.95, 2.6, 0.45, "运维提交完工→待验收\n通知甲方")
    arr(4.1, 6.5, 4.1, 6.2)
    node(4.1, 5.1, 1.8, 0.7, "甲方是否\n确认验收", "decision")
    arr(4.1, 5.7, 4.1, 5.5)
    node(6.4, 5.1, 2.2, 0.45, "退回/继续沟通")
    arr(5.0, 5.1, 5.25, 5.1, "否")
    arr(6.4, 5.35, 4.1, 6.5)
    node(4.1, 4.15, 2.6, 0.45, "工单已完成\n设备恢复正常")
    arr(4.1, 4.7, 4.1, 4.4, "是")
    node(4.1, 3.35, 1.8, 0.7, "甲方是否\n评价", "decision")
    arr(4.1, 3.9, 4.1, 3.75)
    node(4.1, 2.45, 2.6, 0.45, "填写评价→已评价")
    arr(4.1, 2.95, 4.1, 2.7, "是")
    node(4.1, 1.65, 2.6, 0.45, "管理员归档")
    arr(4.1, 2.2, 4.1, 1.9)
    arr(5.05, 3.35, 6.5, 1.65, "否/稍后")
    arr(6.5, 1.65, 5.45, 1.65)
    node(4.1, 0.85, 0, 0, "", "end")
    arr(4.1, 1.4, 4.1, 1.0)
    ax.set_title("图3.0 运维工单报修—验收业务活动图", fontsize=12, pad=6)
    return savefig(fig, "fig3_0_activity.png")


def fig_activity_maintain():
    fig, ax = plt.subplots(figsize=(8.0, 7.5))
    ax.set_xlim(0, 8)
    ax.set_ylim(0, 7.5)
    ax.axis("off")

    def box(x, y, text):
        b = FancyBboxPatch((x - 1.4, y - 0.28), 2.8, 0.56,
                           boxstyle="round,pad=0.02,rounding_size=0.04",
                           facecolor="#F6FFED", edgecolor="#389E0D", lw=1.1)
        ax.add_patch(b)
        ax.text(x, y, text, ha="center", va="center", fontsize=8)

    def arr(y1, y2):
        ax.annotate("", xy=(4, y2), xytext=(4, y1),
                    arrowprops=dict(arrowstyle="->", color="#333", lw=1.1))

    ax.add_patch(plt.Circle((4, 7.0), 0.12, color="#222"))
    box(4, 6.3, "管理员制定维保计划/检查项")
    arr(6.85, 6.6)
    box(4, 5.5, "生成维保任务并指派运维")
    arr(5.98, 5.8)
    box(4, 4.7, "运维按检查项现场执行")
    arr(5.18, 5.0)
    box(4, 3.9, "填报结果/异常/备件并完成")
    arr(4.38, 4.2)
    # decision
    d = plt.Polygon([(4, 3.35), (5.1, 2.85), (4, 2.35), (2.9, 2.85)],
                    closed=True, facecolor="#FFF7E6", edgecolor="#D48806", lw=1.2)
    ax.add_patch(d)
    ax.text(4, 2.85, "是否异常", ha="center", va="center", fontsize=8)
    arr(3.58, 3.2)
    box(6.3, 2.85, "生成关注/转故障处置")
    ax.annotate("", xy=(4.9, 2.85), xytext=(5.15, 2.85),
                arrowprops=dict(arrowstyle="-", color="#333", lw=1.0))
    ax.text(5.0, 3.05, "是", fontsize=7, color="#666")
    box(4, 1.6, "汇总维保档案/更新提醒")
    ax.annotate("", xy=(4, 1.9), xytext=(4, 2.35),
                arrowprops=dict(arrowstyle="->", color="#333", lw=1.1))
    ax.text(4.25, 2.1, "否", fontsize=7, color="#666")
    ax.add_patch(plt.Circle((4, 0.7), 0.14, fill=False, ec="#222", lw=2))
    ax.add_patch(plt.Circle((4, 0.7), 0.07, color="#222"))
    ax.annotate("", xy=(4, 0.88), xytext=(4, 1.28),
                arrowprops=dict(arrowstyle="->", color="#333", lw=1.1))
    ax.set_title("图3.11 预防性维保业务活动图", fontsize=12, pad=6)
    return savefig(fig, "fig3_11_maintain.png")


def generate_all_figs():
    paths = {}
    paths["3_0"] = fig_activity_order()
    paths["3_11"] = fig_activity_maintain()
    paths["3_1"] = fig_role_usecase()
    paths["3_2"] = fig_actor_usecase(
        "图3.2 甲方用户用例图", "fig3_2_client.png", "甲方用户",
        ["登录/退出", "浏览设备台账", "提交报修工单", "查看本人工单",
         "确认验收", "服务评价", "个人中心", "消息中心"])
    paths["3_3"] = fig_actor_usecase(
        "图3.3 浏览设备台账用例图", "fig3_3_device_view.png", "甲方用户",
        ["条件查询设备", "查看设备详情", "选择报修设备"])
    paths["3_4"] = fig_actor_usecase(
        "图3.4 提交报修工单用例图", "fig3_4_submit.png", "甲方用户",
        ["填写故障描述", "上传现场图片", "提交生成工单", "接收进度通知"])
    paths["3_5"] = fig_actor_usecase(
        "图3.5 运维人员用例图", "fig3_5_ops.png", "运维人员",
        ["运维工作台", "处理指派工单", "故障记录", "案例库检索",
         "维保任务执行", "备件领用", "设备健康查看", "消息中心"])
    paths["3_6"] = fig_actor_usecase(
        "图3.6 处理运维工单用例图", "fig3_6_process.png", "运维人员",
        ["接单", "填写处理过程", "登记备件出库", "上传维修图片", "提交完工"])
    paths["3_7"] = fig_actor_usecase(
        "图3.7 管理员用例图", "fig3_7_admin.png", "管理员",
        ["用户权限管理", "冷库冷藏间管理", "设备台账/导入导出", "工单派单管控",
         "维保计划任务", "备件供应商", "健康分析统计", "操作日志/消息"])
    paths["3_8"] = fig_actor_usecase(
        "图3.8 管理用户用例图", "fig3_8_user_mgmt.png", "管理员",
        ["查询用户", "新增/编辑用户", "角色分配", "启用/禁用", "绑定冷藏间"])
    paths["3_9"] = fig_actor_usecase(
        "图3.9 设备台账管理用例图", "fig3_9_device_mgmt.png", "管理员",
        ["设备增删改查", "状态变更", "查看状态历史", "Excel导入", "Excel导出"])
    paths["3_10"] = fig_actor_usecase(
        "图3.10 工单全流程管控用例图", "fig3_10_order_admin.png", "管理员",
        ["查看全量工单", "派单", "转派", "撤回", "关闭/重开", "归档"])
    return paths


# ---------- word helpers ----------
def set_run_font(run, name_cn="宋体", name_en="Times New Roman", size=12, bold=False):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name_en
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name_cn)


def add_para(doc, text, *, style=None, size=12, bold=False, align="left",
             first_indent=True, space_after=6, cn="宋体"):
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.5
    if first_indent and align == "left":
        pf.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_run_font(run, name_cn=cn, size=size, bold=bold)
    return p


def add_heading_cn(doc, text, level=1):
    # custom headings for better CN fonts
    sizes = {1: 16, 2: 14, 3: 12}
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if level > 1 else WD_ALIGN_PARAGRAPH.CENTER
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(12 if level == 1 else 8)
    pf.space_after = Pt(10 if level == 1 else 6)
    pf.line_spacing = 1.5
    run = p.add_run(text)
    set_run_font(run, name_cn="黑体", size=sizes.get(level, 12), bold=True)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    set_run_font(run, name_cn="宋体", size=10.5, bold=False)


def add_image(doc, path, width_cm=14.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))


def add_page_break(doc):
    doc.add_page_break()


def set_page(doc):
    sec = doc.sections[0]
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.17)
    sec.right_margin = Cm(3.17)


def build_doc(figs):
    doc = Document()
    set_page(doc)

    # cover
    for _ in range(3):
        add_para(doc, "", first_indent=False, space_after=0)
    add_para(doc, "本科毕业论文（设计）", size=22, bold=True, align="center",
             first_indent=False, cn="黑体", space_after=18)
    add_para(doc, "基于 Spring Boot + Vue 的冷库设备", size=18, bold=True, align="center",
             first_indent=False, cn="黑体", space_after=4)
    add_para(doc, "综合运维管理系统的设计与实现", size=18, bold=True, align="center",
             first_indent=False, cn="黑体", space_after=28)
    add_para(doc, "（论文前三章：绪论 · 关键技术 · 系统分析）", size=12, align="center",
             first_indent=False, space_after=36)
    add_para(doc, "专业：软件工程", size=14, align="center", first_indent=False, space_after=6)
    add_para(doc, "题目方向：软件应用与开发类", size=14, align="center", first_indent=False, space_after=6)
    add_para(doc, "文档说明：本章内容满足“分析、设计、实现、测试”中分析部分最低要求",
             size=10.5, align="center", first_indent=False, space_after=4)
    add_para(doc, "（含活动图、用例图及用例文字说明）", size=10.5, align="center",
             first_indent=False, space_after=4)

    add_page_break(doc)

    # abstract
    add_heading_cn(doc, "摘　要", 1)
    add_para(doc,
             "针对冷库现场设备种类多、故障处置链条长、维保与库存协同不足等问题，本文设计并实现了一套基于 Spring Boot 与 Vue 的冷库设备综合运维管理系统。"
             "系统面向管理员、运维人员与甲方用户三类角色，覆盖冷库与冷藏间管理、设备台账与生命周期、七状态运维工单闭环、故障记录与案例库、维保计划与任务、"
             "备件与供应商、设备健康分析、数据统计中心以及消息提醒等功能。后端采用 Spring Boot、MyBatis-Plus、MySQL，并结合 Redis、WebSocket 与文件/Excel 能力；"
             "前端采用 Vue2、ElementUI、ECharts 实现业务交互与可视化分析。本文前三章给出研究背景与意义、关键技术及系统需求分析，重点通过活动图与用例图明确业务边界，"
             "为后续系统设计、实现与测试奠定基础。")
    add_para(doc, "关键词：冷库运维；工单管理；Spring Boot；Vue；需求分析", bold=True,
             first_indent=False, space_after=18)

    add_heading_cn(doc, "Abstract", 1)
    add_para(doc,
             "To address multi-device cold-storage sites, long fault-handling chains, and weak collaboration among maintenance and spare parts,"
             " this thesis designs and implements a cold-storage equipment O&M management system based on Spring Boot and Vue."
             " The system serves administrators, maintenance staff and client users, covering storage/room management, device ledger and lifecycle,"
             " a seven-state work-order closed loop, fault records and case library, maintenance plans/tasks, spare parts and suppliers,"
             " rule-based health analysis, statistics center and messaging. The backend uses Spring Boot, MyBatis-Plus and MySQL with Redis/WebSocket/Excel support;"
             " the frontend uses Vue2, ElementUI and ECharts. Chapters 1–3 present background, key technologies and requirements analysis with activity and use-case models.",
             first_indent=False)
    add_para(doc, "Keywords: cold storage O&M; work order; Spring Boot; Vue; requirements analysis",
             bold=True, first_indent=False)

    add_page_break(doc)

    # TOC simple
    add_heading_cn(doc, "目　录", 1)
    toc = [
        "第1章　绪论",
        "　　1.1　研究背景及意义",
        "　　　　1.1.1　研究背景",
        "　　　　1.1.2　研究意义",
        "　　1.2　国内外研究现状",
        "　　　　1.2.1　国外研究现状",
        "　　　　1.2.2　国内研究现状",
        "　　1.3　研究目标及内容",
        "　　　　1.3.1　研究目标",
        "　　　　1.3.2　研究内容",
        "第2章　关键技术介绍",
        "　　2.1　Spring Boot 框架",
        "　　2.2　Vue2 与 ElementUI",
        "　　2.3　MySQL 数据库",
        "　　2.4　MyBatis-Plus",
        "　　2.5　Redis",
        "　　2.6　WebSocket 与文件/Excel 处理",
        "　　2.7　ECharts",
        "　　2.8　本章小结",
        "第3章　系统分析",
        "　　3.1　系统分析概述",
        "　　3.2　系统功能性需求分析",
        "　　　　3.2.1　甲方用户的功能分析",
        "　　　　3.2.2　运维人员的功能分析",
        "　　　　3.2.3　管理员的功能分析",
        "　　3.3　系统非功能性需求分析",
        "　　3.4　本章小结",
    ]
    for t in toc:
        add_para(doc, t, first_indent=False, space_after=2, size=12)

    add_page_break(doc)

    # ===== Chapter 1 =====
    add_heading_cn(doc, "第1章　绪论", 1)
    add_heading_cn(doc, "1.1　研究背景及意义", 2)
    add_heading_cn(doc, "1.1.1　研究背景", 3)
    add_para(doc,
             "冷链仓储是保障生鲜农产品、医药与食品质量安全的重要基础设施。冷库现场通常部署压缩机主机、蒸发器/风机、融霜与排水电热、传感器及控制器等多类设备，"
             "设备一旦异常，可能导致库温波动甚至货物损失。传统运维模式依赖电话报修、纸质台账与经验处置，存在报修信息不完整、派工不透明、维修过程难追溯、"
             "维保计划难落地、备件库存与维修脱节等问题。")
    add_para(doc,
             "随着企业信息化与“设备全生命周期管理”理念普及，冷库运维需要一套面向多角色协同的业务系统：甲方可快速报修并跟踪进度，运维可按指派闭环处理并沉淀经验，"
             "管理员可完成派单管控、台账治理、预防性维保与统计分析。基于 Web 的前后端分离架构开发成本可控、部署灵活，适合作为毕业设计落地的软件应用与开发类作品。")

    add_heading_cn(doc, "1.1.2　研究意义", 3)
    add_para(doc,
             "本研究设计并实现“冷库设备综合运维管理系统”，意义主要体现在：第一，规范故障报修到验收评价的全流程，提升处置时效与过程留痕能力；"
             "第二，将预防性维保与故障维修分开管理，降低重复故障风险；第三，打通工单领用与备件库存，减少缺料延误；第四，通过规则化健康分析与统计看板辅助管理决策；"
             "第五，以软件工程方法完成需求分析、设计、实现与测试，满足软件应用与开发类毕设对模型要素的基本要求。")

    add_heading_cn(doc, "1.2　国内外研究现状", 2)
    add_heading_cn(doc, "1.2.1　国外研究现状", 3)
    add_para(doc,
             "国外在设施运维（Facility Management）与计算机化维护管理系统（CMMS）方面起步较早，常见研究聚焦设备资产台账、预防性维护计划、工单派发与备件库存一体化。"
             "部分冷链相关工作更强调温湿度监测与物联网采集闭环。总体来看，成熟商业系统功能完备，但本地化业务规则、中小冷库场景适配与毕设可实现规模之间存在差距，"
             "直接移植成本较高。")
    add_heading_cn(doc, "1.2.2　国内研究现状", 3)
    add_para(doc,
             "国内冷库信息化研究多集中在组态监控、温度报警或仓储业务系统，针对“报修—派单—维修—验收—评价—归档”完整工单链，以及维保、案例库、健康分析与消息协同的综合运维平台仍相对分散。"
             "部分高校与企业课题已采用 Spring Boot、Vue 技术栈实现仓储或设备管理原型，但在角色权限细化、状态机完整性与库存联动方面仍有提升空间。本文在吸收 CMMS 思想基础上，"
             "面向冷库现场设备对象，构建可落地的多角色运维管理系统。")

    add_heading_cn(doc, "1.3　研究目标及内容", 2)
    add_heading_cn(doc, "1.3.1　研究目标", 3)
    add_para(doc,
             "研究目标是设计并实现一套功能完整、权限清晰、易用性强的冷库设备综合运维管理系统。系统围绕管理员、运维人员、甲方用户构建功能体系，"
             "实现设备台账与生命周期管理、七状态工单闭环、故障与案例沉淀、维保计划任务、备件供应商、健康分析、统计中心与消息提醒，形成可演示、可测试的软件作品。")

    add_heading_cn(doc, "1.3.2　研究内容", 3)
    add_para(doc, "研究内容主要包括：", first_indent=True)
    add_para(doc, "（1）系统需求分析：梳理三角色业务场景，绘制活动图与用例图，撰写用例说明与非功能需求。", first_indent=False)
    add_para(doc, "（2）系统设计：开展总体架构、功能结构、数据库与核心流程设计（后续章节）。", first_indent=False)
    add_para(doc, "（3）系统实现：基于 Spring Boot + Vue2 完成前后端功能开发，集成 Redis、WebSocket、文件上传与 Excel。", first_indent=False)
    add_para(doc, "（4）系统测试：设计功能测试用例验证工单、维保、库存等核心链路（后续章节）。", first_indent=False)

    add_page_break(doc)

    # ===== Chapter 2 =====
    add_heading_cn(doc, "第2章　关键技术介绍", 1)
    add_heading_cn(doc, "2.1　Spring Boot 框架", 2)
    add_para(doc,
             "Spring Boot 是简化 Spring 应用搭建与部署的 Java 开发框架，具备自动配置与 Starter 依赖机制，可快速集成 Web、数据访问等组件，并内嵌 Tomcat 便于独立运行。"
             "本系统后端基于 Spring Boot 构建 RESTful 接口，结合拦截器完成 JWT 鉴权与角色隔离，统一返回结果结构，提升开发效率与可维护性。")

    add_heading_cn(doc, "2.2　Vue2 与 ElementUI", 2)
    add_para(doc,
             "Vue2 是主流前端框架，采用组件化与数据驱动视图的思想，配合 Vue Router 实现路由守卫与页面组织，Axios 负责 HTTP 请求。"
             "ElementUI 提供表格、表单、对话框等成熟组件，适合管理信息系统界面快速构建。本系统前端以 Vue2 + ElementUI 实现多角色菜单与业务页面。")

    add_heading_cn(doc, "2.3　MySQL 数据库", 2)
    add_para(doc,
             "MySQL 是开源关系型数据库，适合存储用户、冷库、设备、工单、库存等结构化业务数据。本系统通过主键/外键与业务编号约束保障数据一致性，"
             "并结合索引与分页查询满足列表检索性能需求。")

    add_heading_cn(doc, "2.4　MyBatis-Plus", 2)
    add_para(doc,
             "MyBatis-Plus 在 MyBatis 基础上提供通用 Mapper、条件构造器与分页插件，显著减少单表 CRUD 样板代码。"
             "本系统使用 MyBatis-Plus 完成实体映射与常见查询，复杂业务在 Service 层封装事务与状态机规则。")

    add_heading_cn(doc, "2.5　Redis", 2)
    add_para(doc,
             "Redis 是内存型键值数据库，常用于缓存与会话辅助。本系统将 Redis 用于登录 Token 黑名单与会话标记等场景；当 Redis 不可用时自动降级为本地缓存，"
             "避免接口因连接超时而阻塞，兼顾毕设环境可用性与技术完整性。")

    add_heading_cn(doc, "2.6　WebSocket 与文件/Excel 处理", 2)
    add_para(doc,
             "WebSocket 支持服务端主动推送，本系统用于消息中心实时提醒。文件上传用于报修/维修图片附件；Apache POI 实现设备台账 Excel 导入导出，"
             "满足批量建档与台账移交等实际需求。")

    add_heading_cn(doc, "2.7　ECharts", 2)
    add_para(doc,
             "ECharts 是常用可视化库，本系统在数据统计中心用于展示工单趋势、故障分布、库存与维保完成情况等图表，支撑管理员从全局视角分析运行状况。")

    add_heading_cn(doc, "2.8　本章小结", 2)
    add_para(doc,
             "本章介绍了系统开发涉及的关键技术：Spring Boot 与 MyBatis-Plus 支撑后端业务与数据访问，Vue2/ElementUI/ECharts 支撑前端交互与可视化，"
             "MySQL 与 Redis 分别承担持久化与缓存辅助，WebSocket 与文件/Excel 能力完善消息与台账场景。上述技术为后续需求分析与系统实现提供支撑。")

    add_page_break(doc)

    # ===== Chapter 3 =====
    add_heading_cn(doc, "第3章　系统分析", 1)
    add_para(doc,
             "按照软件应用与开发类毕设对“分析”部分的要求，本章通过业务需求分析提炼系统需求，给出业务活动图、用例图以及用例文字说明，并补充非功能性需求。",
             first_indent=True)

    add_heading_cn(doc, "3.1　系统分析概述", 2)
    add_para(doc,
             "在冷库设备综合运维管理系统设计开发过程中，围绕设备台账、故障报修与维修闭环、预防性维保、备件库存协同等核心业务，将使用者归纳为管理员、运维人员与甲方用户三类角色，核心需求如下：")
    add_para(doc,
             "（1）甲方用户层面：提供设备查看与故障报修入口，支持上传现场图片；可跟踪本人工单进度，完成待验收确认与服务评价；维护个人信息并接收相关消息。")
    add_para(doc,
             "（2）运维人员层面：处理指派工单（接单、过程填报、备件领用、提交完工）；维护故障记录与检索案例库；执行维保任务；参与备件业务；使用工作台与消息提醒提升响应效率。")
    add_para(doc,
             "（3）管理员层面：负责用户权限、冷库/设备台账、工单派单管控、维保计划、供应商与备件、健康分析与全局统计、操作日志审计，以及导入导出等支撑能力。")

    add_para(doc, "报修—派单—处理—验收—评价—归档的主业务控制流如图3.0所示。", first_indent=True)
    add_image(doc, figs["3_0"], 12.5)
    add_caption(doc, "图3.0　运维工单报修—验收业务活动图")

    add_para(doc, "预防性维保与故障维修相互独立，其业务流程如图3.11所示。", first_indent=True)
    add_image(doc, figs["3_11"], 11.5)
    add_caption(doc, "图3.11　预防性维保业务活动图")

    add_para(doc, "针对系统不同用户角色，给出角色关系用例图，如图3.1所示。", first_indent=True)
    add_image(doc, figs["3_1"], 15)
    add_caption(doc, "图3.1　角色关系用例图")

    add_heading_cn(doc, "3.2　系统功能性需求分析", 2)
    add_heading_cn(doc, "3.2.1　甲方用户的功能分析", 3)
    add_para(doc,
             "甲方用户是冷库现场使用方的业务发起者之一，可进行设备浏览与报修、本人工单跟踪与验收评价、个人中心维护以及消息查看等操作。甲方用户用例图如图3.2所示。")
    add_image(doc, figs["3_2"], 14)
    add_caption(doc, "图3.2　甲方用户用例图")

    add_para(doc, "（1）登录与退出", bold=True, first_indent=False)
    add_para(doc,
             "甲方通过用户名密码登录，系统校验账号状态并签发 JWT；登录后按角色展示菜单。退出时清除本地会话并将 Token 写入黑名单，防止会话被复用。")

    add_para(doc, "（2）浏览设备台账", bold=True, first_indent=False)
    add_para(doc,
             "甲方可按编号、名称、冷藏间、类型、状态查询设备并查看详情。甲方不可增删改设备，但可查看全部设备（含公共设备），以便准确选择报修对象。用例图如图3.3所示。")
    add_image(doc, figs["3_3"], 13)
    add_caption(doc, "图3.3　浏览设备台账用例图")

    # use case table for submit
    add_para(doc, "（3）提交报修工单", bold=True, first_indent=False)
    add_para(doc,
             "提交报修是甲方核心用例。甲方选择设备、填写故障描述并可上传图片，系统生成唯一工单号，初始状态为“待受理”，联动设备故障状态，并向管理员推送新报修消息。用例图如图3.4所示。")
    add_image(doc, figs["3_4"], 13)
    add_caption(doc, "图3.4　提交报修工单用例图")

    # case description table
    add_para(doc, "表3.1　“提交报修工单”用例说明", bold=True, align="center", first_indent=False, size=10.5)
    table = doc.add_table(rows=8, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rows = [
        ("用例名称", "提交报修工单"),
        ("参与者", "甲方用户"),
        ("前置条件", "甲方已登录；系统中存在可选设备"),
        ("后置条件", "生成待受理工单；设备状态联动；管理员收到消息"),
        ("基本事件流", "1.进入工单/报修页面；2.选择设备；3.填写描述并可选上传图片；4.提交；5.系统生成工单号并提示成功"),
        ("异常事件流", "未选择设备/描述为空时提示校验失败；上传超限文件时拒绝"),
        ("业务规则", "甲方仅可操作本人提交的工单；允许报修公共设备与普通设备"),
        ("优先级", "高"),
    ]
    for i, (a, b) in enumerate(rows):
        table.rows[i].cells[0].text = a
        table.rows[i].cells[1].text = b
        for cell in table.rows[i].cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    set_run_font(r, size=10.5)

    add_para(doc, "（4）查看本人工单进度", bold=True, first_indent=False, space_after=4)
    add_para(doc, "甲方仅可查看本人提交的工单列表与详情，包括状态、指派运维、处理过程、维修图片与备件等信息。")

    add_para(doc, "（5）确认验收与服务评价", bold=True, first_indent=False, space_after=4)
    add_para(doc,
             "工单处于“待验收”时，甲方确认完工后进入“已完成”，设备恢复正常，并通知运维。对已完成工单可进行满意度评价，工单进入“已评价”。")

    add_para(doc, "（6）维护个人信息与消息中心", bold=True, first_indent=False, space_after=4)
    add_para(doc,
             "甲方可修改姓名、联系方式与密码，查看绑定冷藏间归属展示信息；可在消息中心查看待验收等提醒，支持已读标记与 WebSocket 实时角标。")

    add_heading_cn(doc, "3.2.2　运维人员的功能分析", 3)
    add_para(doc,
             "运维人员是故障处置与预防性维保的执行主体，可处理指派工单、维护故障与案例、执行维保任务、参与备件业务，并使用工作台与消息提醒。运维人员用例图如图3.5所示。")
    add_image(doc, figs["3_5"], 14)
    add_caption(doc, "图3.5　运维人员用例图")

    add_para(doc, "（1）运维工作台", bold=True, first_indent=False, space_after=4)
    add_para(doc, "展示个人待处理工单数、本月完成数与临近维保清单，支持快速跳转处置。")

    add_para(doc, "（2）处理运维工单", bold=True, first_indent=False, space_after=4)
    add_para(doc,
             "对“已派单”工单接单进入“处理中”；填写过程、原因、方案、耗时，可上传维修图片并登记备件出库；提交完工后进入“待验收”。"
             "列表可看较广范围，但处理类操作仅限指派给本人的工单。用例图如图3.6所示。")
    add_image(doc, figs["3_6"], 13)
    add_caption(doc, "图3.6　处理运维工单用例图")

    add_para(doc, "表3.2　“处理运维工单”用例说明", bold=True, align="center", first_indent=False, size=10.5)
    table2 = doc.add_table(rows=7, cols=2)
    table2.style = "Table Grid"
    rows2 = [
        ("用例名称", "处理运维工单"),
        ("参与者", "运维人员"),
        ("前置条件", "运维已登录；存在指派给本人的已派单/处理中工单"),
        ("后置条件", "工单进入待验收；库存按领用扣减；甲方收到待验收消息"),
        ("基本事件流", "1.查看指派工单；2.接单；3.填报处理信息与备件；4.上传维修图片；5.提交完工"),
        ("异常事件流", "处理非本人工单被拒绝；备件库存不足时出库失败并提示"),
        ("优先级", "高"),
    ]
    for i, (a, b) in enumerate(rows2):
        table2.rows[i].cells[0].text = a
        table2.rows[i].cells[1].text = b
        for cell in table2.rows[i].cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    set_run_font(r, size=10.5)

    add_para(doc, "（3）故障记录与案例库", bold=True, first_indent=False, space_after=4)
    add_para(doc,
             "运维可维护故障记录并关联设备/工单；可按设备类型、故障类型与关键词检索案例库获取排查与处理参考（非 AI）。")

    add_para(doc, "（4）维保任务、备件与消息", bold=True, first_indent=False, space_after=4)
    add_para(doc,
             "运维执行指派维保任务并填报检查结果；查询备件与流水，在领用场景触发出库；接收派单/维保指派等实时消息。")

    add_heading_cn(doc, "3.2.3　管理员的功能分析", 3)
    add_para(doc,
             "管理员负责平台配置与全流程管控，覆盖用户权限、基础台账、工单派单、维保计划、库存治理、健康与统计、日志审计等。管理员用例图如图3.7所示。")
    add_image(doc, figs["3_7"], 14)
    add_caption(doc, "图3.7　管理员用例图")

    add_para(doc, "（1）用户与权限管理", bold=True, first_indent=False, space_after=4)
    add_para(doc, "支持用户增删改查、角色分配、启用禁用及甲方冷藏间绑定。用例图如图3.8所示。")
    add_image(doc, figs["3_8"], 13)
    add_caption(doc, "图3.8　管理用户用例图")

    add_para(doc, "（2）冷库/冷藏间与设备台账", bold=True, first_indent=False, space_after=4)
    add_para(doc,
             "维护冷库与冷藏间档案；库区详情定位为单库办事台（待办工单、异常设备、临近维保）。设备台账支持生命周期状态历史与 Excel 导入导出。用例图如图3.9所示。")
    add_image(doc, figs["3_9"], 13)
    add_caption(doc, "图3.9　设备台账管理用例图")

    add_para(doc, "（3）工单全流程管控", bold=True, first_indent=False, space_after=4)
    add_para(doc,
             "管理员可查看全平台工单并执行派单、转派、撤回、关闭、重开与归档。状态机为：待受理→已派单→处理中→待验收→已完成→已评价→已归档。用例图如图3.10所示。")
    add_image(doc, figs["3_10"], 13)
    add_caption(doc, "图3.10　工单全流程管控用例图")

    add_para(doc, "表3.3　“工单派单”用例说明", bold=True, align="center", first_indent=False, size=10.5)
    table3 = doc.add_table(rows=6, cols=2)
    table3.style = "Table Grid"
    rows3 = [
        ("用例名称", "工单派单"),
        ("参与者", "管理员"),
        ("前置条件", "存在待受理工单；系统中有可用运维账号"),
        ("后置条件", "工单变为已派单；运维收到派单消息"),
        ("基本事件流", "1.打开待受理工单；2.选择运维人员；3.确认派单；4.系统更新状态并推送消息"),
        ("优先级", "高"),
    ]
    for i, (a, b) in enumerate(rows3):
        table3.rows[i].cells[0].text = a
        table3.rows[i].cells[1].text = b
        for cell in table3.rows[i].cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    set_run_font(r, size=10.5)

    add_para(doc, "（4）维保、库存、健康分析与统计", bold=True, first_indent=False, space_after=4)
    add_para(doc,
             "管理员制定维保计划并生成/指派任务；维护供应商与备件，设置安全库存并接收不足预警；按规则计算设备健康等级；在统计中心查看工单、故障、维保与库存多维图表。")

    add_heading_cn(doc, "3.3　系统非功能性需求分析", 2)
    add_heading_cn(doc, "3.3.1　可用性", 3)
    add_para(doc,
             "面向三角色提供统一 Web 界面，核心流程步骤清晰；列表支持多条件查询与分页；统计页图表化展示；消息角标与实时推送降低漏看待办概率。")
    add_heading_cn(doc, "3.3.2　安全性", 3)
    add_para(doc,
             "密码 BCrypt 加密；JWT 鉴权与拦截器校验；退出后 Token 黑名单；按角色隔离菜单与接口；甲方仅本人单据，运维仅处理本人任务；关键操作写日志。")
    add_heading_cn(doc, "3.3.3　可靠性", 3)
    add_para(doc,
             "工单状态机、库存扣减与设备状态联动在服务端校验并使用事务；Redis 故障自动降级，避免接口长时间阻塞。")
    add_heading_cn(doc, "3.3.4　可扩展性", 3)
    add_para(doc,
             "前后端分离与模块化划分，便于扩展设备类型、维保模板、报表维度或通知通道。")
    add_heading_cn(doc, "3.3.5　可维护性", 3)
    add_para(doc,
             "代码按业务模块组织，配置集中管理，数据库变更以迁移脚本维护，便于部署与问题定位。")

    add_heading_cn(doc, "3.4　本章小结", 2)
    add_para(doc,
             "本章完成了冷库设备综合运维管理系统的需求分析：通过活动图描述报修闭环与维保流程，通过用例图与用例说明明确三角色功能边界，并从可用性、安全性、可靠性、可扩展性与可维护性给出非功能需求。"
             "分析结果满足软件应用与开发类作品对分析阶段“活动图、用例图、用例说明”等模型要素要求，为后续系统设计与实现提供依据。")

    doc.save(OUT)
    return OUT


def main():
    figs = generate_all_figs()
    out = build_doc(figs)
    print("OK", out)
    for k, v in figs.items():
        print(k, v)


if __name__ == "__main__":
    main()
