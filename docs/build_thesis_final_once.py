# -*- coding: utf-8 -*-
"""
按周思汉论文用例图风格重绘（扇形关联 + 框内 2x2/分栏 + <<include>>），
保证椭圆在框内、标签不压线；全部完成后只生成 1 份 Word。
"""
from pathlib import Path
import math
import shutil
import matplotlib.pyplot as plt
from matplotlib.patches import Ellipse, Rectangle
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

ROOT = Path(r"f:\plcmin")
FIG = ROOT / "docs" / "thesis_figs"
FIG.mkdir(parents=True, exist_ok=True)
# 唯一最终输出
FINAL_DOCX = ROOT / "docs" / "冷库设备运维管理系统-论文第1-4章.docx"

plt.rcParams["font.sans-serif"] = ["SimSun", "Microsoft YaHei", "SimHei"]
plt.rcParams["axes.unicode_minus"] = False


def save(fig, name):
    p = FIG / name
    fig.savefig(p, dpi=200, bbox_inches="tight", facecolor="white", pad_inches=0.2)
    plt.close(fig)
    return p


# 线宽对齐周思汉图3.2：细线
LW_ACTOR = 0.5
LW_OVAL = 0.55
LW_FRAME = 0.6
LW_ASSOC = 0.45
LW_INCLUDE = 0.45


def actor(ax, x, y, label):
    ax.add_patch(plt.Circle((x, y), 0.12, fill=False, ec="black", lw=LW_ACTOR, zorder=5))
    ax.plot([x, x], [y - 0.12, y - 0.48], "k-", lw=LW_ACTOR, zorder=5)
    ax.plot([x - 0.18, x + 0.18], [y - 0.26, y - 0.26], "k-", lw=LW_ACTOR, zorder=5)
    ax.plot([x, x - 0.14], [y - 0.48, y - 0.72], "k-", lw=LW_ACTOR, zorder=5)
    ax.plot([x, x + 0.14], [y - 0.48, y - 0.72], "k-", lw=LW_ACTOR, zorder=5)
    ax.text(x, y - 0.92, label, ha="center", va="top", fontsize=12)


def oval(ax, x, y, text, w=2.35, h=0.78):
    ax.add_patch(Ellipse((x, y), w, h, facecolor="white", edgecolor="black", lw=LW_OVAL, zorder=4))
    ax.text(x, y, text, ha="center", va="center", fontsize=11, zorder=6)
    return x, y, w, h


def box(ax, x0, y0, x1, y1, title):
    ax.add_patch(Rectangle((x0, y0), x1 - x0, y1 - y0, fill=False, ec="black", lw=LW_FRAME, zorder=1))
    # 标题在框内顶部，下方留空给椭圆
    ax.text((x0 + x1) / 2, y1 - 0.32, title, ha="center", va="top", fontsize=13)


def fan(ax, ax_xy, targets):
    """从参与者中心扇形连到各椭圆左缘（范本风格，非强制水平梳齿）。"""
    ax0, ay0 = ax_xy
    for tx, ty, tw, th in targets:
        # 连接到椭圆左端点
        ex, ey = tx - tw / 2.0, ty
        ax.plot([ax0 + 0.2, ex], [ay0 - 0.15, ey], "k-", lw=LW_ASSOC, zorder=2)


def include(ax, a, b, side=0.0, t=0.5):
    """
    <<include>> 标在虚线上（对齐周思汉图3.2）：
    - 白底不透明，只截断虚线，不压椭圆
    - t∈(0,1) 控制标签在线段上的位置，避免多条 include 挤在同一点
    """
    x1, y1, w1, h1 = a
    x2, y2, w2, h2 = b
    if abs(x2 - x1) >= abs(y2 - y1) * 0.6:
        sx, sy = x1 + w1 / 2.0, y1
        ex, ey = x2 - w2 / 2.0, y2
    else:
        sx, sy = x1, y1 - h1 / 2.0
        ex, ey = x2, y2 + h2 / 2.0
    ax.annotate(
        "",
        xy=(ex, ey),
        xytext=(sx, sy),
        arrowprops=dict(arrowstyle="->", color="black", lw=LW_INCLUDE, linestyle=(0, (4, 2.5))),
        zorder=3,
    )
    mx = sx + (ex - sx) * t
    my = sy + (ey - sy) * t
    if side:
        dx, dy = ex - sx, ey - sy
        L = math.hypot(dx, dy) or 1.0
        mx += -dy / L * side
        my += dx / L * side
    ax.text(
        mx,
        my,
        "<<include>>",
        ha="center",
        va="center",
        fontsize=8,
        zorder=8,
        # 不透明白底：只盖住虚线，不透字缝
        bbox=dict(boxstyle="square,pad=0.08", facecolor="white", edgecolor="none", alpha=1.0),
    )


def style_detail(fname, title, actor_name, left_top, left_bot, right_top, right_bot="登录"):
    """
    严格仿范本「浏览课程列表 / 浏览设备台账」：
    参与者 --扇形--> 左上、左下；左上/左下 <<include>> 右上；右上 <<include>> 右下登录。
    列距/行距加大，三条 include 标签错开，互不遮挡。
    """
    fig, ax = plt.subplots(figsize=(9.6, 5.9))
    ax.set_xlim(0, 9.6)
    ax.set_ylim(0, 5.9)
    ax.axis("off")

    x0, y0, x1, y1 = 2.35, 0.4, 9.35, 5.55
    box(ax, x0, y0, x1, y1, title)

    # 加大列距、行距，给虚线标签留空
    lt = oval(ax, 4.35, 4.05, left_top, w=2.25, h=0.78)
    lb = oval(ax, 4.35, 2.05, left_bot, w=2.25, h=0.78)
    rt = oval(ax, 7.55, 4.05, right_top, w=2.25, h=0.78)
    rb = oval(ax, 7.55, 1.95, right_bot, w=2.05, h=0.72)

    actor(ax, 1.05, 3.1, actor_name)
    fan(ax, (1.05, 3.1), [lt, lb])

    # 上横：标签偏左段；斜线：标签偏下段；竖线：居中 —— 三处错开
    include(ax, lt, rt, t=0.42)
    include(ax, lb, rt, t=0.38)
    include(ax, rt, rb, t=0.55)

    return save(fig, fname)


def style_detail3(fname, title, actor_name, lefts, right_top, right_bot="登录"):
    """左列 3 个 + 右上 + 登录。"""
    fig, ax = plt.subplots(figsize=(9.2, 6.2))
    ax.set_xlim(0, 9.2)
    ax.set_ylim(0, 6.2)
    ax.axis("off")
    box(ax, 2.4, 0.4, 8.95, 5.85, title)

    ys = [4.7, 3.35, 2.0]
    left_ovals = [oval(ax, 4.45, y, name, w=2.3, h=0.75) for y, name in zip(ys, lefts)]
    rt = oval(ax, 7.2, 4.7, right_top, w=2.3, h=0.75)
    rb = oval(ax, 7.2, 2.5, right_bot, w=2.1, h=0.75)

    actor(ax, 1.05, 3.35, actor_name)
    fan(ax, (1.05, 3.35), left_ovals)

    for o in left_ovals:
        include(ax, o, rt)
    include(ax, rt, rb)
    return save(fig, fname)


def style_overview(fname, title, actor_name, cases):
    """
    角色总览：左侧单列主用例（参与者扇形连接，仿范本），
    右侧仅「登录」；只保留 2 条 include，标签落在中间空白，避免压椭圆文字。
    """
    mains = [c for c in cases if c != "登录"]
    n = len(mains)
    row_h = 0.88
    title_band = 0.95
    height = 1.05 + title_band + n * row_h + 0.55
    fig, ax = plt.subplots(figsize=(9.6, max(5.6, height * 0.82)))
    ax.set_xlim(0, 9.6)
    ax.set_ylim(0, height)
    ax.axis("off")

    x0, y0, x1, y1 = 2.45, 0.4, 9.35, height - 0.22
    box(ax, x0, y0, x1, y1, title)

    left_x = 4.55
    w = 2.35
    first_y = y1 - title_band - 0.12
    ovals = []
    for i, name in enumerate(mains):
        y = first_y - i * row_h
        ovals.append(oval(ax, left_x, y, name, w=w, h=0.7))

    # 登录：右侧垂直居中偏下，与左列拉开间距
    login_y = (ovals[0][1] + ovals[-1][1]) / 2 - 0.15
    login = oval(ax, 7.55, login_y, "登录", w=2.15, h=0.72)

    ay = (ovals[0][1] + ovals[-1][1]) / 2
    actor(ax, 1.1, ay, actor_name)
    fan(ax, (1.1, ay), ovals)

    # include 标签压在虚线中点上
    include(ax, ovals[0], login)
    include(ax, ovals[-1], login)

    return save(fig, fname)


def style_role():
    """
    三角色关系图：范本风格系统框 + 用例；
    每位参与者扇形连到自己的用例子集，避免总线；用例单列靠右，留出连线区。
    """
    groups = [
        ("管理员", ["用户与权限管理", "工单派单管控", "健康分析与统计", "操作日志审计"]),
        ("运维人员", ["处理运维工单", "维保任务执行", "故障案例备件", "备件领用"]),
        ("甲方用户", ["浏览设备台账", "提交报修工单", "确认验收评价", "消息中心"]),
    ]
    row_h = 0.78
    title_band = 1.0
    gap = 0.35
    n = sum(len(g[1]) for g in groups)
    height = 0.9 + title_band + n * row_h + 2 * gap + 0.5
    fig, ax = plt.subplots(figsize=(10.2, height * 0.82))
    ax.set_xlim(0, 10.2)
    ax.set_ylim(0, height)
    ax.axis("off")

    x0, y0, x1, y1 = 3.4, 0.4, 9.95, height - 0.22
    box(ax, x0, y0, x1, y1, "冷库设备综合运维管理系统")

    uc_x = (x0 + x1) / 2 + 0.35
    y = y1 - title_band - 0.1
    for actor_name, cases in groups:
        ovals = []
        for name in cases:
            ovals.append(oval(ax, uc_x, y, name, w=2.85, h=0.62))
            y -= row_h
        ay = (ovals[0][1] + ovals[-1][1]) / 2
        actor(ax, 1.2, ay, actor_name)
        fan(ax, (1.2, ay), ovals)
        y -= gap

    return save(fig, "fig3_1_role.png")


def redraw_all_usecases():
    style_role()
    style_overview(
        "fig3_2_client.png",
        "冷库设备运维管理系统",
        "甲方用户",
        ["浏览设备台账", "提交报修工单", "查看本人工单", "确认验收", "服务评价", "个人中心", "消息中心", "登录"],
    )
    style_overview(
        "fig3_5_ops.png",
        "冷库设备运维管理系统",
        "运维人员",
        ["运维工作台", "处理指派工单", "故障记录", "案例库检索", "维保任务执行", "备件领用", "设备健康查看", "消息中心"],
    )
    style_overview(
        "fig3_7_admin.png",
        "冷库设备运维管理系统",
        "管理员",
        ["用户权限管理", "冷库冷藏间管理", "设备台账管理", "工单派单管控", "维保计划任务", "备件供应商", "健康分析统计", "操作日志"],
    )

    # 分项：范本 2x2（两关联 + include），线宽同上
    style_detail("fig3_3_device_view.png", "浏览设备台账", "甲方用户", "条件查询设备", "选择报修设备", "查看设备详情")
    style_detail("fig3_4_submit.png", "提交报修工单", "甲方用户", "填写故障描述", "上传现场图片", "提交生成工单")
    style_detail("fig3_6_process.png", "处理运维工单", "运维人员", "接单", "提交完工", "填写过程与备件")
    style_detail("fig3_8_user_mgmt.png", "管理用户", "管理员", "新增用户", "启用/禁用", "查询用户")
    style_detail("fig3_9_device_mgmt.png", "设备台账管理", "管理员", "设备增删改查", "状态变更", "Excel导入导出")
    style_detail("fig3_10_order_admin.png", "工单全流程管控", "管理员", "派单", "转派", "关闭/归档")
    print("usecase figs done")


# ---------- 活动图等若缺失则简要保留已有；Word 组装 ----------
def set_run_font(run, name_cn="宋体", name_en="Times New Roman", size=12, bold=False):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name_en
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name_cn)


def add_para(doc, text, *, size=12, bold=False, align="left", first_indent=True, space_after=6, cn="宋体"):
    p = doc.add_paragraph()
    p.alignment = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }.get(align, WD_ALIGN_PARAGRAPH.JUSTIFY)
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.5
    if first_indent and align == "left":
        pf.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_run_font(run, name_cn=cn, size=size, bold=bold)


def add_h(doc, text, level=1):
    sizes = {1: 16, 2: 14, 3: 12}
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(10 if level == 1 else 6)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    set_run_font(run, name_cn="黑体", size=sizes.get(level, 12), bold=True)


def add_cap(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    set_run_font(p.add_run(text), size=10.5)


def add_img(doc, path, w=13.0):
    if not Path(path).exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(w))


def add_table(doc, title, headers, rows):
    add_para(doc, title, bold=True, align="center", first_indent=False, size=10.5)
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        table.rows[0].cells[j].text = h
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            table.rows[i + 1].cells[j].text = str(val)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    set_run_font(r, size=9)
    add_para(doc, "", first_indent=False, space_after=4)


def ensure_extra_figs():
    """活动图/架构图若已存在则复用；缺失时用占位提示跳过。"""
    needed = [
        "fig3_0_activity.png",
        "fig3_11_maintain.png",
        "fig4_1_arch.png",
        "fig4_2_func.png",
        "fig4_3_user_flow.png",
        "fig4_4_admin_flow.png",
        "fig4_5_er.png",
    ]
    missing = [n for n in needed if not (FIG / n).exists()]
    if missing:
        print("warn missing figs (will skip images):", missing)


def cleanup_old_variants():
    """清理此前多次生成的中间 docx，只保留最终一份。"""
    keep = {FINAL_DOCX.name}
    for p in (ROOT / "docs").glob("冷库设备运维管理系统-论文*.docx"):
        if p.name not in keep:
            try:
                p.unlink()
                print("removed", p.name)
            except Exception as e:
                print("skip remove", p.name, e)


def build_one_docx():
    """一次性写出完整第1–4章（含已重绘用例图）。"""
    ensure_extra_figs()
    f = lambda n: FIG / n

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.17)
    sec.right_margin = Cm(3.17)

    for _ in range(2):
        add_para(doc, "", first_indent=False, space_after=0)
    add_para(doc, "本科毕业论文（设计）", size=22, bold=True, align="center", first_indent=False, cn="黑体", space_after=16)
    add_para(doc, "基于 Spring Boot + Vue 的冷库设备", size=18, bold=True, align="center", first_indent=False, cn="黑体", space_after=4)
    add_para(doc, "综合运维管理系统的设计与实现", size=18, bold=True, align="center", first_indent=False, cn="黑体", space_after=20)
    add_para(doc, "（第1–4章：绪论 · 关键技术 · 系统分析 · 系统设计）", size=12, align="center", first_indent=False, space_after=28)
    add_para(doc, "专业：软件工程", size=14, align="center", first_indent=False, space_after=6)
    add_para(doc, "题目方向：软件应用与开发类", size=14, align="center", first_indent=False, space_after=6)

    doc.add_page_break()
    add_h(doc, "摘　要", 1)
    add_para(
        doc,
        "针对冷库现场设备种类多、故障处置链条长、维保与库存协同不足等问题，本文设计并实现了一套基于 Spring Boot 与 Vue 的冷库设备综合运维管理系统。"
        "系统面向管理员、运维人员与甲方用户三类角色，覆盖冷库与冷藏间管理、设备台账与生命周期、七状态运维工单闭环、故障记录与案例库、维保计划与任务、"
        "备件与供应商、设备健康分析、数据统计中心以及消息提醒等功能。后端采用 Spring Boot、MyBatis-Plus、MySQL，并结合 Redis、WebSocket 与文件/Excel 能力；"
        "前端采用 Vue2、ElementUI、ECharts。本文前四章给出研究背景、关键技术、需求分析与系统设计。",
    )
    add_para(doc, "关键词：冷库运维；工单管理；Spring Boot；Vue；需求分析；系统设计", bold=True, first_indent=False)

    doc.add_page_break()
    add_h(doc, "目　录", 1)
    for t in [
        "第1章　绪论",
        "　　1.1　研究背景及意义",
        "　　1.2　国内外研究现状",
        "　　1.3　研究目标及内容",
        "第2章　关键技术介绍",
        "第3章　系统分析",
        "　　3.1　系统分析概述",
        "　　3.2　系统功能性需求分析",
        "　　3.3　系统非功能性需求分析",
        "　　3.4　本章小结",
        "第4章　系统设计",
        "　　4.1　系统技术架构设计",
        "　　4.2　系统功能设计",
        "　　4.3　数据库设计",
        "　　4.4　本章小结",
    ]:
        add_para(doc, t, first_indent=False, space_after=2)

    # Ch1
    doc.add_page_break()
    add_h(doc, "第1章　绪论", 1)
    add_h(doc, "1.1　研究背景及意义", 2)
    add_h(doc, "1.1.1　研究背景", 3)
    add_para(
        doc,
        "冷链仓储是保障生鲜农产品、医药与食品质量安全的重要基础设施。冷库现场通常部署压缩机主机、蒸发器/风机、融霜与排水电热、传感器及控制器等多类设备，"
        "设备异常可能导致库温波动甚至货损。传统运维依赖电话报修与纸质台账，存在信息不完整、派工不透明、过程难追溯、维保难落地、备件与维修脱节等问题。",
    )
    add_para(
        doc,
        "企业信息化与设备全生命周期管理理念普及后，需要一套多角色协同系统：甲方快速报修与跟踪，运维按指派闭环处置并沉淀经验，管理员完成派单、台账、维保与统计。"
        "基于 Web 的前后端分离架构适合作为软件应用与开发类毕业设计落地。",
    )
    add_h(doc, "1.1.2　研究意义", 3)
    add_para(
        doc,
        "意义主要体现在：规范报修到验收评价全流程；将预防性维保与故障维修分轨管理；打通工单领用与备件库存；通过规则化健康分析与统计看板辅助决策；"
        "并以软件工程方法完成分析、设计、实现与测试。",
    )
    add_h(doc, "1.2　国内外研究现状", 2)
    add_h(doc, "1.2.1　国外研究现状", 3)
    add_para(
        doc,
        "国外 CMMS/设施运维研究较成熟，强调资产台账、预防性维护、工单与备件一体化；冷链方向常结合监测采集。商业系统完备，但本地化规则与毕设可实现规模存在差距。",
    )
    add_h(doc, "1.2.2　国内研究现状", 3)
    add_para(
        doc,
        "国内研究多集中在组态监控或仓储业务，完整工单链及维保、案例、健康分析、消息协同的综合运维平台仍相对分散。"
        "本文吸收 CMMS 思想，面向冷库设备对象构建可落地的多角色运维系统。",
    )
    add_h(doc, "1.3　研究目标及内容", 2)
    add_h(doc, "1.3.1　研究目标", 3)
    add_para(
        doc,
        "设计并实现功能完整、权限清晰、易用的冷库设备综合运维管理系统，形成可演示、可测试的软件作品。",
    )
    add_h(doc, "1.3.2　研究内容", 3)
    add_para(doc, "（1）需求分析：活动图、用例图与用例说明、非功能需求。", first_indent=False)
    add_para(doc, "（2）系统设计：架构、功能结构、时序/类设计、数据库设计。", first_indent=False)
    add_para(doc, "（3）系统实现与测试（后续章节）。", first_indent=False)

    # Ch2
    doc.add_page_break()
    add_h(doc, "第2章　关键技术介绍", 1)
    for title, body in [
        ("2.1　Spring Boot 框架", "Spring Boot 简化 Spring 应用搭建与部署，具备自动配置与 Starter 机制。本系统后端据此构建 REST 接口，并结合拦截器完成 JWT 鉴权。"),
        ("2.2　Vue2 与 ElementUI", "Vue2 组件化开发，配合 Vue Router 与 Axios；ElementUI 提供管理端常用组件，支撑多角色业务页面。"),
        ("2.3　MySQL 数据库", "MySQL 存储用户、冷库、设备、工单、库存等结构化数据，通过约束与索引保障一致性与检索性能。"),
        ("2.4　MyBatis-Plus", "提供通用 Mapper 与条件构造器，减少单表 CRUD 样板代码，复杂规则在 Service 层封装。"),
        ("2.5　Redis", "用于 Token 黑名单等缓存场景；不可用时降级本地缓存，避免接口长时间阻塞。"),
        ("2.6　WebSocket 与文件/Excel", "WebSocket 推送消息提醒；本地上传保存图片；POI 实现设备台账导入导出。"),
        ("2.7　ECharts", "用于统计中心工单、故障、维保、库存等可视化图表。"),
        ("2.8　本章小结", "上述技术分别支撑后端服务、前端交互、数据持久化、缓存推送与可视化，为系统分析与设计提供基础。"),
    ]:
        add_h(doc, title, 2)
        add_para(doc, body)

    # Ch3
    doc.add_page_break()
    add_h(doc, "第3章　系统分析", 1)
    add_para(
        doc,
        "按照软件应用与开发类毕设对分析部分的要求，本章给出业务活动图、用例图及用例文字说明，并补充非功能性需求。用例图绘制参考软件工程范本风格：参与者扇形关联、系统边界内椭圆用例，以及必要的 <<include>> 关系，并保证图形元素互不遮挡。",
    )
    add_h(doc, "3.1　系统分析概述", 2)
    add_para(
        doc,
        "系统使用者归纳为管理员、运维人员与甲方用户。（1）甲方：设备浏览与报修、本人工单跟踪、验收评价、消息与个人中心；"
        "（2）运维：处理指派工单、故障与案例、维保任务、备件领用、工作台与消息；"
        "（3）管理员：用户权限、台账、派单管控、维保计划、库存、健康与统计、日志审计。",
    )
    add_para(doc, "报修—派单—处理—验收主业务活动如图3.0所示。", first_indent=True)
    add_img(doc, f("fig3_0_activity.png"), 12.5)
    add_cap(doc, "图3.0　运维工单报修—验收业务活动图")
    add_para(doc, "预防性维保流程如图3.11所示。", first_indent=True)
    add_img(doc, f("fig3_11_maintain.png"), 11.5)
    add_cap(doc, "图3.11　预防性维保业务活动图")
    add_para(doc, "角色关系用例图如图3.1所示。", first_indent=True)
    add_img(doc, f("fig3_1_role.png"), 13.5)
    add_cap(doc, "图3.1　角色关系用例图")

    add_h(doc, "3.2　系统功能性需求分析", 2)
    add_h(doc, "3.2.1　甲方用户的功能分析", 3)
    add_para(doc, "甲方用户可进行设备浏览与报修、本人工单跟踪与验收评价、个人中心与消息查看等。甲方用户用例图如图3.2所示。")
    add_img(doc, f("fig3_2_client.png"), 13.2)
    add_cap(doc, "图3.2　甲方用户用例图")

    add_para(doc, "（1）登录与退出", bold=True, first_indent=False)
    add_para(doc, "甲方通过账号密码登录，系统签发 JWT；退出时清除会话并将 Token 写入黑名单。")
    add_para(doc, "（2）浏览设备台账", bold=True, first_indent=False)
    add_para(doc, "甲方可查询并查看全部设备（含公共设备），不可增删改。浏览设备台账用例图如图3.3所示。")
    add_img(doc, f("fig3_3_device_view.png"), 13.0)
    add_cap(doc, "图3.3　浏览设备台账用例图")

    add_para(doc, "（3）提交报修工单", bold=True, first_indent=False)
    add_para(doc, "甲方选择设备填写故障描述并可上传图片，系统生成待受理工单并通知管理员。用例图如图3.4所示。")
    add_img(doc, f("fig3_4_submit.png"), 13.0)
    add_cap(doc, "图3.4　提交报修工单用例图")

    add_para(doc, "表3.1　“提交报修工单”用例说明", bold=True, align="center", first_indent=False, size=10.5)
    add_table(
        doc,
        "",
        ["项目", "内容"],
        [
            ["用例名称", "提交报修工单"],
            ["参与者", "甲方用户"],
            ["前置条件", "已登录；存在可选设备"],
            ["后置条件", "生成待受理工单；设备状态联动；管理员收到消息"],
            ["基本事件流", "选择设备→填写描述/上传图片→提交→生成工单号"],
            ["异常事件流", "校验失败或上传超限时提示并拒绝"],
            ["优先级", "高"],
        ],
    )

    add_para(doc, "（4）查看本人工单、确认验收与评价、个人中心与消息中心", bold=True, first_indent=False)
    add_para(doc, "甲方仅查看本人单据；待验收可确认完工；完成后可评价；支持消息已读与实时提醒。")

    add_h(doc, "3.2.2　运维人员的功能分析", 3)
    add_para(doc, "运维人员处理指派工单、维护故障与案例、执行维保、参与备件业务，并使用工作台与消息。用例图如图3.5所示。")
    add_img(doc, f("fig3_5_ops.png"), 13.2)
    add_cap(doc, "图3.5　运维人员用例图")
    add_para(doc, "处理运维工单是核心用例：接单、填报过程与备件、上传维修图片、提交完工。用例图如图3.6所示。")
    add_img(doc, f("fig3_6_process.png"), 13.0)
    add_cap(doc, "图3.6　处理运维工单用例图")
    add_table(
        doc,
        "表3.2　“处理运维工单”用例说明",
        ["项目", "内容"],
        [
            ["用例名称", "处理运维工单"],
            ["参与者", "运维人员"],
            ["前置条件", "存在指派给本人的已派单/处理中工单"],
            ["后置条件", "工单待验收；库存按领用扣减；甲方收到提醒"],
            ["基本事件流", "接单→填报→登记备件→上传图片→提交完工"],
            ["优先级", "高"],
        ],
    )

    add_h(doc, "3.2.3　管理员的功能分析", 3)
    add_para(doc, "管理员负责平台配置与全流程管控。用例图如图3.7所示。")
    add_img(doc, f("fig3_7_admin.png"), 13.2)
    add_cap(doc, "图3.7　管理员用例图")
    add_para(doc, "用户管理用例如图3.8所示；设备台账管理如图3.9所示；工单全流程管控如图3.10所示。")
    add_img(doc, f("fig3_8_user_mgmt.png"), 13.0)
    add_cap(doc, "图3.8　管理用户用例图")
    add_img(doc, f("fig3_9_device_mgmt.png"), 13.0)
    add_cap(doc, "图3.9　设备台账管理用例图")
    add_img(doc, f("fig3_10_order_admin.png"), 13.0)
    add_cap(doc, "图3.10　工单全流程管控用例图")
    add_table(
        doc,
        "表3.3　“工单派单”用例说明",
        ["项目", "内容"],
        [
            ["用例名称", "工单派单"],
            ["参与者", "管理员"],
            ["前置条件", "存在待受理工单与可用运维账号"],
            ["后置条件", "工单已派单；运维收到消息"],
            ["基本事件流", "选择运维→确认派单→更新状态并推送"],
            ["优先级", "高"],
        ],
    )

    add_h(doc, "3.3　系统非功能性需求分析", 2)
    for t, b in [
        ("3.3.1　可用性", "统一 Web 界面与清晰流程；列表筛选分页；统计图表化；消息角标与推送降低漏看待办。"),
        ("3.3.2　安全性", "BCrypt 密码；JWT 鉴权；退出黑名单；三角色菜单/接口隔离；关键操作写日志。"),
        ("3.3.3　可靠性", "状态机与库存事务校验；Redis 故障自动降级。"),
        ("3.3.4　可扩展性", "前后端分离与模块化，便于扩展字典、模板与通知通道。"),
        ("3.3.5　可维护性", "按业务模块组织代码，配置集中，库表变更可迁移。"),
    ]:
        add_h(doc, t, 3)
        add_para(doc, b)

    add_h(doc, "3.4　本章小结", 2)
    add_para(
        doc,
        "本章完成需求分析：活动图描述主流程，用例图按范本风格绘制并保证不遮挡，用例说明与非功能需求为后续设计提供依据。",
    )

    # Ch4（结构对齐周思汉：4.1架构/功能图 → 4.2用户端/管理端流程 → 4.3库表 → 图号4.1~4.5）
    doc.add_page_break()
    add_h(doc, "第4章　系统设计", 1)
    add_para(
        doc,
        "在前述需求分析基础上，本章对冷库设备综合运维管理系统进行系统设计，包括技术架构、功能结构、用户端与管理端流程设计以及数据库设计。"
        "设计内容与第3章用例保持一致。",
    )
    add_h(doc, "4.1　系统技术架构设计", 2)
    add_h(doc, "4.1.1　系统软件层次架构", 3)
    add_para(
        doc,
        "本系统整体采用前后端分离架构。前端为 Vue2 单页应用，按管理员、运维人员、甲方用户三类角色组织页面；"
        "后端基于 Spring Boot，按控制层、业务逻辑层、数据访问层与数据层分层实现，并结合 JWT 鉴权、Redis 可降级缓存与 WebSocket 消息推送。"
        "系统架构如图4.1所示。",
    )
    add_img(doc, f("fig4_1_arch.png"), 14.2)
    add_cap(doc, "图4.1　系统架构图")
    add_h(doc, "4.1.2　系统功能结构", 3)
    add_para(
        doc,
        "系统功能按“用户端—管理端”划分。用户端面向甲方用户与运维人员；管理端由管理员负责用户权限、台账、工单派单、维保备件、统计与日志等。"
        "系统功能图如图4.2所示。",
    )
    add_img(doc, f("fig4_2_func.png"), 14.5)
    add_cap(doc, "图4.2　系统功能图")

    add_h(doc, "4.2　系统的功能设计", 2)
    add_para(doc, "系统主要分为用户端与管理端两大功能模块，下面分别进行功能设计说明。")
    add_h(doc, "4.2.1　用户端", 3)
    add_para(
        doc,
        "用户端面向甲方用户与运维人员。用户登录后，系统根据角色区分权限与菜单。"
        "甲方用户可浏览设备、提交报修、验收评价并查看消息；运维人员可接单处理、填报过程与备件、提交完工，并执行维保与故障案例维护。"
        "用户端流程图如图4.3所示。",
    )
    add_img(doc, f("fig4_3_user_flow.png"), 13.5)
    add_cap(doc, "图4.3　用户端流程图")
    add_h(doc, "4.2.2　管理端", 3)
    add_para(
        doc,
        "管理端面向系统管理员。管理员登录后可查看数据概览，完成用户管理、冷库设备台账、工单派单管控、维保备件、统计日志及个人信息维护。"
        "管理端流程图如图4.4所示。",
    )
    add_img(doc, f("fig4_4_admin_flow.png"), 13.5)
    add_cap(doc, "图4.4　管理端流程图")

    add_h(doc, "4.3　数据库设计", 2)
    add_para(
        doc,
        "在技术架构与功能设计基础上进行数据库设计。业务数据以 MySQL 结构化存储为主，Redis 用于 Token 黑名单等可降级缓存，上传文件存于本地 uploads 目录。",
    )
    add_h(doc, "4.3.1　数据库的概念设计", 3)
    add_para(
        doc,
        "概念模型围绕用户、冷库—冷藏间—设备、运维工单、故障与维保、备件与供应商、消息等实体展开。"
        "系统 E-R 图如图4.5所示。",
    )
    add_img(doc, f("fig4_5_er.png"), 14.2)
    add_cap(doc, "图4.5　冷库设备运维管理系统E-R图")
    add_h(doc, "4.3.2　数据库的逻辑设计", 3)
    add_para(doc, "系统共 19 张业务表。以下选取核心表说明结构，其余表与之关联支撑完整业务。")
    add_table(doc, "表4.1　用户表 sys_user", ["字段", "类型", "说明"], [
        ["id", "BIGINT", "主键"], ["username", "VARCHAR", "登录名"], ["password", "VARCHAR", "BCrypt"],
        ["role", "VARCHAR", "ADMIN/OPS/CLIENT"], ["status", "TINYINT", "启停"],
    ])
    add_table(doc, "表4.2　冷库表 cold_storage", ["字段", "类型", "说明"], [
        ["id", "BIGINT", "主键"], ["name", "VARCHAR", "名称"], ["address", "VARCHAR", "地址"], ["status", "VARCHAR", "状态"],
    ])
    add_table(doc, "表4.3　冷藏间表 cold_room", ["字段", "类型", "说明"], [
        ["id", "BIGINT", "主键"], ["storage_id", "BIGINT", "冷库FK"], ["name/code", "VARCHAR", "名称/编码"],
    ])
    add_table(doc, "表4.4　设备表 device", ["字段", "类型", "说明"], [
        ["id", "BIGINT", "主键"], ["device_no", "VARCHAR", "编号"], ["room_id", "BIGINT", "冷藏间FK"],
        ["status", "VARCHAR", "生命周期状态"], ["is_public", "TINYINT", "公共设备"],
    ])
    add_table(doc, "表4.5　工单表 work_order", ["字段", "类型", "说明"], [
        ["id", "BIGINT", "主键"], ["order_no", "VARCHAR", "工单号"], ["device_id", "BIGINT", "设备FK"],
        ["status", "VARCHAR", "七状态"], ["submit_user_id/assignee_id", "BIGINT", "报修人/运维"],
    ])
    add_table(doc, "表4.6　备件表 spare_part", ["字段", "类型", "说明"], [
        ["id", "BIGINT", "主键"], ["part_no", "VARCHAR", "编号"], ["stock_qty", "INT", "库存"],
        ["safety_stock", "INT", "安全库存"], ["supplier_id", "BIGINT", "供应商FK"],
    ])
    add_table(doc, "表4.7　消息表 sys_message", ["字段", "类型", "说明"], [
        ["id", "BIGINT", "主键"], ["user_id", "BIGINT", "接收人"], ["msg_type", "VARCHAR", "类型"], ["is_read", "TINYINT", "已读"],
    ])
    add_para(
        doc,
        "其余表：user_cold_room、operation_log、device_status_log、work_order_part、fault_record、fault_case、"
        "maintain_plan、maintain_plan_item、maintain_task、maintain_task_item、supplier、stock_record，分别支撑绑定、审计、状态史、工单备件、故障案例、维保与库存流水。",
        first_indent=True,
    )

    add_h(doc, "4.4　本章小结", 2)
    add_para(
        doc,
        "本章完成系统设计：给出层次架构图与系统功能图，按用户端与管理端分别给出流程图，并完成数据库概念结构与核心表逻辑设计，"
        "与第3章需求对应，为后续系统实现提供依据。",
    )

    # 先清理旧变体，再保存唯一文件
    cleanup_old_variants()
    # 若 FINAL 被占用，写临时名再提示
    try:
        doc.save(str(FINAL_DOCX))
        print("FINAL", FINAL_DOCX)
    except PermissionError:
        alt = ROOT / "docs" / "冷库设备运维管理系统-论文第1-4章-请关闭旧文件后看这个.docx"
        doc.save(str(alt))
        print("FINAL_LOCKED_ALT", alt)


def main():
    redraw_all_usecases()
    build_one_docx()


if __name__ == "__main__":
    main()
